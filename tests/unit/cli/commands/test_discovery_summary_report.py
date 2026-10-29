"""Tests for `discovery summary-report`.

The command's own decisions are which records make the report and where it is written; the
workbook itself is covered by the builder's tests.
"""

import datetime
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from statgpt.cli.commands import discovery as discovery_module
from statgpt.cli.commands.discovery import summary_report_handler
from statgpt.cli.shared.discovery_report import SUMMARIES_SHEET
from statgpt.common.schemas import (
    Channel,
    DiscoveryDataset,
    DiscoveryIndexingStatus,
    DiscoverySummaryCall,
    DiscoverySummaryStep,
    DiscoveryValidationStatus,
)

_NOW = datetime.datetime(2026, 1, 1, tzinfo=datetime.timezone.utc)
_CHANNEL_ID = 7
_DEPLOYMENT_ID = "my-channel"


def _channel() -> Channel:
    return Channel(
        id=_CHANNEL_ID,
        created_at=_NOW,
        updated_at=_NOW,
        title="My Channel",
        description="",
        deployment_id=_DEPLOYMENT_ID,
        llm_model="gpt-4o",
        details={  # type: ignore[arg-type]
            "supremeAgent": {"name": "bot", "domain": "stats", "terminologyDomain": "stats"}
        },
    )


def _record(dataset_id: str, **kwargs: Any) -> DiscoveryDataset:
    fields: dict[str, Any] = {
        "id": 1,
        "created_at": _NOW,
        "updated_at": _NOW,
        "channel_id": _CHANNEL_ID,
        "agency": "IMF",
        "dataset_id": dataset_id,
        "validation_status": DiscoveryValidationStatus.VALID,
        "indexing_status": DiscoveryIndexingStatus.INDEXED,
    }
    return DiscoveryDataset(**(fields | kwargs))


_FAILED_CALL = DiscoverySummaryCall(step=DiscoverySummaryStep.CHUNK, part=1, error="cut short")

_RECORDS = [
    _record("SUMMARIZED", indicators_summarized_at=_NOW),
    _record("FAILED", indicators_summary_failed_calls=[_FAILED_CALL]),
    _record("COPIED"),
]


class _StubAdminClient:
    def __init__(self, records: list[DiscoveryDataset]) -> None:
        self._records = records
        self.requests: list[tuple[int, bool]] = []

    async def health_check(self) -> bool:
        return True

    async def get_channels(self) -> list[Channel]:
        return [_channel()]

    async def get_discovery_datasets(
        self, channel_id: int, *, include_summary_calls: bool = False
    ) -> list[DiscoveryDataset]:
        self.requests.append((channel_id, include_summary_calls))
        return self._records


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch) -> _StubAdminClient:
    stub = _StubAdminClient(list(_RECORDS))

    @asynccontextmanager
    async def get_admin_client(*_: Any, **__: Any):
        yield stub

    monkeypatch.setattr(discovery_module, "get_admin_client", get_admin_client)
    return stub


def _reported_ids(path: Path) -> set[str]:
    sheet = load_workbook(path)[SUMMARIES_SHEET]
    return {str(sheet.cell(row=row, column=2).value) for row in range(2, sheet.max_row + 1)}


async def test_the_report_holds_summarized_and_failed_records(
    client: _StubAdminClient, tmp_path: Path
) -> None:
    output = tmp_path / "report.xlsx"

    await summary_report_handler(channel=_DEPLOYMENT_ID, output=str(output))

    assert client.requests == [(_CHANNEL_ID, True)]
    assert _reported_ids(output) == {"SUMMARIZED", "FAILED"}


async def test_all_includes_records_no_llm_summarized(
    client: _StubAdminClient, tmp_path: Path
) -> None:
    output = tmp_path / "report.xlsx"

    await summary_report_handler(channel=_DEPLOYMENT_ID, output=str(output), all=True)

    assert _reported_ids(output) == {"SUMMARIZED", "FAILED", "COPIED"}


async def test_nothing_summarized_writes_no_file(
    client: _StubAdminClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    client._records = [_record("COPIED")]
    monkeypatch.chdir(tmp_path)

    await summary_report_handler(channel=_DEPLOYMENT_ID)

    assert list(tmp_path.iterdir()) == []


async def test_the_default_file_is_named_after_the_channel(
    client: _StubAdminClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)

    await summary_report_handler(channel=_DEPLOYMENT_ID)

    [written] = tmp_path.iterdir()
    assert written.name.startswith(f"discovery-summaries-{_DEPLOYMENT_ID}-")
    assert written.suffix == ".xlsx"
