"""Tests for the workbook of `discovery summary-report`.

A record's calls live on a sheet of its own, reached by a link, so what matters is that every
record gets a sheet Excel accepts, that the links land where they say, and that no cell holds
what Excel would repair the file over.
"""

import datetime
import io
from typing import Any

from openpyxl import Workbook, load_workbook

from statgpt.cli.shared.discovery_report import (
    MAX_CELL_CHARS,
    SUMMARIES_SHEET,
    build_summary_report,
)
from statgpt.common.schemas import (
    DiscoveryDataset,
    DiscoveryIndexingStatus,
    DiscoveryLLMCallUsage,
    DiscoverySummaryCall,
    DiscoverySummaryRun,
    DiscoverySummaryStep,
    DiscoveryValidationStatus,
)

_NOW = datetime.datetime(2026, 1, 1, tzinfo=datetime.timezone.utc)


def _usage(
    prompt_tokens: int = 1_000,
    completion_tokens: int = 200,
    cost: float | None = 0.01,
    **kwargs: Any,
) -> DiscoveryLLMCallUsage:
    fields: dict[str, Any] = {
        "deployment": "gpt-4o",
        "model": "gpt-4o-2024",
        "duration_seconds": 2.0,
    }
    return DiscoveryLLMCallUsage(
        prompt_tokens=prompt_tokens,
        completion_tokens=completion_tokens,
        cost=cost,
        **(fields | kwargs),
    )


def _call(**kwargs: Any) -> DiscoverySummaryCall:
    kwargs.setdefault("usage", _usage())
    return DiscoverySummaryCall(step=DiscoverySummaryStep.CHUNK, part=1, **kwargs)


def _run(*calls: DiscoverySummaryCall) -> DiscoverySummaryRun:
    return DiscoverySummaryRun(calls=list(calls))


def _accepted() -> DiscoverySummaryCall:
    return _call(answer={"reasoning": "plan", "units": None, "groups": ["GDP, by sector", "CPI"]})


def _record(agency: str = "IMF", dataset_id: str = "WEO", **kwargs: Any) -> DiscoveryDataset:
    fields: dict[str, Any] = {
        "id": 1,
        "created_at": _NOW,
        "updated_at": _NOW,
        "channel_id": 7,
        "agency": agency,
        "dataset_id": dataset_id,
        "indicators_coverage": [{"name": "GDP", "units": "USD"}, {"name": "CPI", "units": "Index"}],
        "indicators_summary": "GDP, by sector; CPI",
        "indicators_summarized_at": _NOW,
        "indicators_summary_calls": _run(_accepted()),
        "validation_status": DiscoveryValidationStatus.VALID,
        "indexing_status": DiscoveryIndexingStatus.INDEXED,
    }
    return DiscoveryDataset(**(fields | kwargs))


def _round_trip(workbook: Workbook) -> Workbook:
    """Save and load, so the test sees what Excel would open."""
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return load_workbook(buffer)


def _reported_ids(workbook: Workbook) -> list[Any]:
    """The ids of the reported records, the totals row left out."""
    sheet = workbook[SUMMARIES_SHEET]
    return [sheet.cell(row=row, column=2).value for row in range(2, sheet.max_row)]


def _rows(sheet: Any, header_row: int) -> list[dict[str, Any]]:
    header = [cell.value for cell in sheet[header_row]]
    return [
        dict(zip(header, (cell.value for cell in row)))
        for row in sheet.iter_rows(min_row=header_row + 1)
    ]


def test_each_record_with_calls_gets_a_sheet_linked_both_ways() -> None:
    workbook = _round_trip(build_summary_report([_record()]))

    link = workbook[SUMMARIES_SHEET]["B2"].hyperlink
    assert link.location == "'WEO'!A1"
    back = workbook["WEO"]["A1"].hyperlink
    assert back.location == f"'{SUMMARIES_SHEET}'!B2"


def test_a_call_row_shows_the_answer() -> None:
    rejected = _call(answer={"reasoning": "first", "groups": ["x"]}, error="- answer: too long")
    record = _record(indicators_summary_calls=_run(rejected, _accepted()))

    sheet = _round_trip(build_summary_report([record]))["WEO"]

    rows = _rows(sheet, 3)
    assert [row["Accepted"] for row in rows] == [False, True]
    assert rows[0]["Error"] == "- answer: too long"
    assert rows[1]["Reasoning"] == "plan"
    assert rows[1]["Groups"] == "GDP, by sector\nCPI"


def test_failed_run_calls_follow_the_succeeded_ones() -> None:
    record = _record(indicators_summary_failed_calls=_run(_call(answer=None, error="cut short")))

    workbook = _round_trip(build_summary_report([record]))

    rows = _rows(workbook["WEO"], 3)
    assert [row["Run"] for row in rows] == ["succeeded", "failed"]
    assert rows[1]["Reasoning"] is None  # an answer that was not JSON leaves the cells empty
    assert _rows(workbook[SUMMARIES_SHEET], 1)[0]["Last run failed"] is True


def test_the_same_id_under_two_agencies_gets_two_sheets() -> None:
    records = [_record("IMF", "WEO"), _record("OECD", "weo")]

    workbook = _round_trip(build_summary_report(records))

    assert workbook.sheetnames == [SUMMARIES_SHEET, "WEO", "weo (2)"]
    assert workbook[SUMMARIES_SHEET]["B3"].hyperlink.location == "'weo (2)'!A1"


def test_ids_no_sheet_may_be_named_after_are_made_safe() -> None:
    """Records stored before ids had to be sheet-safe are still reported."""
    long_id = "A" * 40
    records = [
        _record(dataset_id="a/b:c"),
        _record(dataset_id="Summaries"),
        _record(dataset_id=long_id),
        _record(dataset_id="'quoted'"),
    ]

    workbook = _round_trip(build_summary_report(records))

    assert set(workbook.sheetnames) == {
        SUMMARIES_SHEET,
        "a_b_c",
        "Summaries (2)",
        "A" * 31,
        "quoted",
    }


def test_a_long_id_clash_keeps_within_the_sheet_name_limit() -> None:
    records = [_record("IMF", "A" * 40), _record("OECD", "A" * 40)]

    workbook = _round_trip(build_summary_report(records))

    assert workbook.sheetnames[2] == "A" * 27 + " (2)"


def test_a_record_without_calls_gets_no_sheet() -> None:
    record = _record(indicators_summarized_at=None, indicators_summary_calls=None)

    workbook = _round_trip(build_summary_report([record]))

    assert workbook.sheetnames == [SUMMARIES_SHEET]
    assert workbook[SUMMARIES_SHEET]["B2"].hyperlink is None
    row = _rows(workbook[SUMMARIES_SHEET], 1)[0]
    assert (row["Model"], row["Usage"], row["LLM seconds"]) == (None, None, None)


def test_text_past_the_cell_limit_is_cut_with_a_marker() -> None:
    record = _record(indicators_coverage=[{"name": "x" * (MAX_CELL_CHARS + 10)}])

    workbook = _round_trip(build_summary_report([record]))

    value = str(_rows(workbook[SUMMARIES_SHEET], 1)[0]["Indicators coverage"])
    assert len(value) == MAX_CELL_CHARS
    assert value.endswith("[truncated]")


def test_text_starting_with_equals_is_not_a_formula() -> None:
    record = _record(indicators_summary_calls=_run(_call(answer={"groups": ["=SUM(A1)"]})))

    sheet = _round_trip(build_summary_report([record]))["WEO"]

    assert sheet["L4"].value == "=SUM(A1)"
    assert sheet["L4"].data_type == "s"


def test_records_are_ordered_by_agency_then_id() -> None:
    records = [_record("OECD", "B"), _record("imf", "Z"), _record("IMF", "A")]

    workbook = _round_trip(build_summary_report(records))

    assert _reported_ids(workbook) == ["A", "Z", "B"]


def test_a_call_row_shows_its_usage() -> None:
    usage = _usage(12_345, 1_234, 0.0123, duration_seconds=3.5)
    record = _record(indicators_summary_calls=_run(_call(usage=usage)))

    [row] = _rows(_round_trip(build_summary_report([record]))["WEO"], 3)

    assert row["Model"] == "gpt-4o-2024"
    assert row["Usage"] == "12,345 in + 1,234 out tokens, $0.0123"
    assert row["Seconds"] == 3.5


def test_a_record_row_sums_the_usage_of_both_runs() -> None:
    record = _record(
        indicators_summary_calls=_run(_accepted(), _call(usage=_usage(model="other-model"))),
        indicators_summary_failed_calls=_run(_call(error="cut short")),
    )

    row = _rows(_round_trip(build_summary_report([record]))[SUMMARIES_SHEET], 1)[0]

    assert row["Model"] == "gpt-4o-2024\nother-model"
    assert row["Usage"] == "3,000 in + 600 out tokens, $0.0300"
    assert row["LLM seconds"] == 6.0


def test_an_unknown_cost_is_shown_as_such() -> None:
    record = _record(indicators_summary_calls=_run(_accepted(), _call(usage=_usage(cost=None))))

    workbook = _round_trip(build_summary_report([record]))

    assert _rows(workbook["WEO"], 3)[1]["Usage"] == "1,000 in + 200 out tokens, cost n/a"
    record_row, total = _rows(workbook[SUMMARIES_SHEET], 1)
    assert record_row["Usage"] == "2,000 in + 400 out tokens, cost n/a"
    assert total["Usage"] == "2,000 in + 400 out tokens, cost n/a"


def test_the_totals_row_sums_every_record_outside_the_autofilter() -> None:
    records = [
        _record("IMF", "A", indicators_summary_calls=_run(_call(error="too long"), _accepted())),
        _record("IMF", "B"),
        _record("IMF", "C", indicators_summary_calls=None, indicators_summarized_at=None),
    ]

    sheet = _round_trip(build_summary_report(records))[SUMMARIES_SHEET]

    total = _rows(sheet, 1)[-1]
    assert total["Agency"] == "TOTAL"
    assert (total["Answers"], total["Rejected"]) == (3, 1)
    assert total["Usage"] == "3,000 in + 600 out tokens, $0.0300"
    assert total["LLM seconds"] == 6.0
    assert sheet.auto_filter.ref.endswith(str(sheet.max_row - 1))
