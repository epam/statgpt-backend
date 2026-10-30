"""The Excel report of a channel's discovery indicator summaries, for `discovery summary-report`.

One `Summaries` sheet lists every record, with a totals row, and each record with summarization
calls gets a sheet of its own holding them, linked both ways.
"""

import datetime
from collections.abc import Iterable, Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, ValidationError

from statgpt.common.schemas import (
    DiscoveryDataset,
    DiscoveryLLMCallUsage,
    DiscoveryLLMUsage,
    DiscoverySummaryCall,
    render_flat_indicators,
)
from statgpt.common.schemas.discovery_dataset import (
    FORBIDDEN_DATASET_ID_CHARS,
    MAX_DATASET_ID_LENGTH,
)

SUMMARIES_SHEET = "Summaries"

MAX_CELL_CHARS = 32_767
"""What Excel holds in one cell. openpyxl writes longer text, and Excel then repairs the file."""

_TRUNCATED = "… [truncated]"

_RESERVED_SHEET_NAMES = (SUMMARIES_SHEET, "History")
"""Taken before any dataset sheet is named: the report's own sheet, and the one Excel keeps."""

_FALLBACK_SHEET_NAME = "dataset"

_CALLS_HEADER_ROW = 3

_BOLD = Font(bold=True)
_LINK = Font(color="0563C1", underline="single")
_TOP = Alignment(vertical="top")
_WRAPPED = Alignment(vertical="top", wrap_text=True)

_SUMMARY_COLUMNS = [
    ("Agency", 24, False),
    ("Dataset ID", 20, False),
    ("Name", 40, True),
    ("Validation", 14, False),
    ("Indexing", 12, False),
    ("Summarized at (UTC)", 20, False),
    ("Indicators", 11, False),
    ("Coverage chars", 15, False),
    ("Summary chars", 15, False),
    ("Ratio", 8, False),
    ("Parts", 7, False),
    ("Answers", 9, False),
    ("Rejected", 9, False),
    ("Model", 24, True),
    ("Usage", 30, True),
    ("LLM seconds", 12, False),
    ("Last run failed", 15, False),
    ("Indicators coverage", 60, False),
    ("Indicators summary", 80, True),
]
"""Header, width and whether the text wraps. The raw coverage does not: it can run for pages."""

_RATIO_COLUMN = 10
_SECONDS_FORMAT = "0.0"

_CALL_COLUMNS = [
    ("Run", 10, False),
    ("#", 5, False),
    ("Step", 8, False),
    ("Part", 6, False),
    ("Accepted", 10, False),
    ("Model", 24, True),
    ("Usage", 30, True),
    ("Seconds", 9, False),
    ("Error", 50, True),
    ("Reasoning", 60, True),
    ("Units", 30, True),
    ("Groups", 80, True),
]


class _Answer(BaseModel):
    """The fields of a stored answer the report shows.

    Lenient on purpose: answers are stored as the model returned them, under whichever output
    model the run used, and a report must not fail on an older shape.
    """

    model_config = ConfigDict(extra="ignore")

    reasoning: str | None = None
    units: str | None = None
    groups: list[str] | None = None


def _parse_answer(call: DiscoverySummaryCall) -> _Answer:
    if call.answer is None:
        return _Answer()
    try:
        return _Answer.model_validate(call.answer)
    except ValidationError:
        return _Answer()


def _text(value: str) -> str:
    """Text as a cell can hold it: no control characters, within Excel's cell limit."""
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) > MAX_CELL_CHARS:
        return value[: MAX_CELL_CHARS - len(_TRUNCATED)] + _TRUNCATED
    return value


def _sheet_name(dataset_id: str, taken: set[str]) -> str:
    """A sheet name for a record, unique among `taken` ignoring case, which it is added to.

    The dataset id is not unique on its own: two agencies can publish the same id, and records
    stored before ids had to be sheet-safe can hold characters no sheet name may.
    """
    base = "".join("_" if char in FORBIDDEN_DATASET_ID_CHARS else char for char in dataset_id)
    base = base.strip("'")[:MAX_DATASET_ID_LENGTH].strip("'") or _FALLBACK_SHEET_NAME

    name = base
    number = 1
    while name.casefold() in taken:
        number += 1
        suffix = f" ({number})"
        name = base[: MAX_DATASET_ID_LENGTH - len(suffix)] + suffix

    taken.add(name.casefold())
    return name


def _link(cell: Any, sheet: str, target: str) -> None:
    """Make a cell link to a cell of a sheet in the same workbook."""
    quoted = sheet.replace("'", "''")
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{quoted}'!{target}")
    cell.font = _LINK


def _write_row(sheet: Worksheet, row: int, values: Iterable[Any], wrapped: Sequence[bool]) -> None:
    for column, (value, wrap) in enumerate(zip(values, wrapped), start=1):
        if isinstance(value, str):
            value = _text(value)
        cell = sheet.cell(row=row, column=column, value=value)
        # openpyxl takes text starting with '=' for a formula; an answer is never one.
        if isinstance(value, str) and cell.data_type == "f":
            cell.data_type = "s"
        cell.alignment = _WRAPPED if wrap else _TOP


def _write_header(
    sheet: Worksheet, row: int, columns: Sequence[tuple[str, int, bool]], last_row: int
) -> None:
    for column, (title, width, _) in enumerate(columns, start=1):
        sheet.cell(row=row, column=column, value=title).font = _BOLD
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = f"A{row + 1}"
    sheet.auto_filter.ref = f"A{row}:{get_column_letter(len(columns))}{max(last_row, row)}"


def _excel_datetime(value: datetime.datetime | None) -> datetime.datetime | None:
    """A datetime Excel can store: it has no time zones, so UTC, and naive."""
    if value is None or value.tzinfo is None:
        return value
    return value.astimezone(datetime.timezone.utc).replace(tzinfo=None)


def _calls(record: DiscoveryDataset) -> list[DiscoverySummaryCall]:
    run = record.indicators_summary_calls
    return run.calls if run else []


def _failed_calls(record: DiscoveryDataset) -> list[DiscoverySummaryCall]:
    run = record.indicators_summary_failed_calls
    return run.calls if run else []


def _usage(record: DiscoveryDataset) -> list[DiscoveryLLMUsage]:
    """The usage of both of the record's runs."""
    runs = (record.indicators_summary_calls, record.indicators_summary_failed_calls)
    return [item for run in runs if run for item in run.usage]


def _format_usage(prompt_tokens: int, completion_tokens: int, cost: float | None) -> str:
    cost_text = f"${cost:,.4f}" if cost is not None else "cost n/a"
    return f"{prompt_tokens:,} in + {completion_tokens:,} out tokens, {cost_text}"


def _sum_cost(costs: Iterable[float | None]) -> float | None:
    """The total of the costs; `None` when any of them is unknown."""
    total = 0.0
    for cost in costs:
        if cost is None:
            return None
        total += cost
    return total


def _summed_usage(usage: Sequence[DiscoveryLLMUsage | DiscoveryLLMCallUsage]) -> str | None:
    if not usage:
        return None
    return _format_usage(
        sum(item.prompt_tokens for item in usage),
        sum(item.completion_tokens for item in usage),
        _sum_cost(item.cost for item in usage),
    )


def _summary_values(record: DiscoveryDataset) -> list[Any]:
    calls = _calls(record)
    usage = _usage(record)
    coverage = render_flat_indicators(record.indicators_coverage)
    coverage_chars = len(coverage)
    summary_chars = len(record.indicators_summary)
    return [
        record.agency,
        record.dataset_id,
        record.name,
        record.validation_status.value,
        record.indexing_status.value,
        _excel_datetime(record.indicators_summarized_at),
        len(record.indicators_coverage),
        coverage_chars,
        summary_chars,
        summary_chars / coverage_chars if coverage_chars else None,
        len({call.part for call in calls if call.part is not None}),
        len(calls),
        sum(1 for call in calls if call.error),
        "\n".join(dict.fromkeys(item.model for item in usage)) or None,
        _summed_usage(usage),
        sum(item.duration_seconds for item in usage) if usage else None,
        bool(_failed_calls(record)),
        coverage,
        record.indicators_summary,
    ]


def _call_rows(record: DiscoveryDataset) -> Iterable[list[Any]]:
    runs = [
        ("succeeded", _calls(record)),
        ("failed", _failed_calls(record)),
    ]
    for run, calls in runs:
        for number, call in enumerate(calls, start=1):
            answer = _parse_answer(call)
            yield [
                run,
                number,
                call.step.value,
                call.part,
                call.error is None,
                call.usage.model,
                _summed_usage([call.usage]),
                call.usage.duration_seconds,
                call.error,
                answer.reasoning,
                answer.units,
                "\n".join(answer.groups) if answer.groups else None,
            ]


def _write_calls_sheet(
    sheet: Worksheet, record: DiscoveryDataset, summaries_row: int, title: str
) -> None:
    back = sheet.cell(row=1, column=1, value=f"← {SUMMARIES_SHEET}")
    _link(back, SUMMARIES_SHEET, f"B{summaries_row}")
    sheet.cell(row=1, column=2, value=_text(title)).font = _BOLD

    wrapped = [wrap for _, _, wrap in _CALL_COLUMNS]
    row = _CALLS_HEADER_ROW
    seconds_column = _column(_CALL_COLUMNS, "Seconds")
    for row, values in enumerate(_call_rows(record), start=_CALLS_HEADER_ROW + 1):
        _write_row(sheet, row, values, wrapped)
        sheet.cell(row=row, column=seconds_column).number_format = _SECONDS_FORMAT
    _write_header(sheet, _CALLS_HEADER_ROW, _CALL_COLUMNS, row)


def _has_calls(record: DiscoveryDataset) -> bool:
    return bool(_calls(record) or _failed_calls(record))


def _column(columns: Sequence[tuple[str, int, bool]], title: str) -> int:
    return next(number for number, (name, _, _) in enumerate(columns, start=1) if name == title)


def _write_totals(sheet: Worksheet, row: int, records: Sequence[DiscoveryDataset]) -> None:
    """The channel's totals, below the records and outside their autofilter."""
    calls = [call for record in records for call in _calls(record)]
    usage = [item for record in records for item in _usage(record)]
    totals: dict[str, str | int | float | None] = {
        "Agency": "TOTAL",
        "Answers": len(calls),
        "Rejected": sum(1 for call in calls if call.error),
        "Usage": _summed_usage(usage),
        "LLM seconds": sum(item.duration_seconds for item in usage),
    }
    for title, value in totals.items():
        column = _column(_SUMMARY_COLUMNS, title)
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = _BOLD
        cell.alignment = _WRAPPED if _SUMMARY_COLUMNS[column - 1][2] else _TOP
    sheet.cell(row=row, column=_column(_SUMMARY_COLUMNS, "LLM seconds")).number_format = (
        _SECONDS_FORMAT
    )


def build_summary_report(records: Sequence[DiscoveryDataset]) -> Workbook:
    """Lay out the records' indicator summaries and their LLM calls as a workbook."""
    workbook = Workbook()
    summaries = workbook.worksheets[0]
    summaries.title = SUMMARIES_SHEET

    taken = {name.casefold() for name in _RESERVED_SHEET_NAMES}
    wrapped = [wrap for _, _, wrap in _SUMMARY_COLUMNS]
    ordered = sorted(records, key=lambda r: (r.agency.casefold(), r.dataset_id.casefold()))

    seconds_column = _column(_SUMMARY_COLUMNS, "LLM seconds")
    row = 1
    for row, record in enumerate(ordered, start=2):
        _write_row(summaries, row, _summary_values(record), wrapped)
        summaries.cell(row=row, column=_RATIO_COLUMN).number_format = "0%"
        summaries.cell(row=row, column=seconds_column).number_format = _SECONDS_FORMAT
        if _has_calls(record):
            name = _sheet_name(record.dataset_id, taken)
            _write_calls_sheet(
                workbook.create_sheet(name), record, row, f"{record.agency} · {record.dataset_id}"
            )
            _link(summaries.cell(row=row, column=2), name, "A1")
    _write_header(summaries, 1, _SUMMARY_COLUMNS, row)
    _write_totals(summaries, row + 1, ordered)

    return workbook
